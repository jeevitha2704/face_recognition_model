import os
try:
    import cv2
except ImportError:
    cv2 = None
import pickle
import sqlite3
import numpy as np
from flask import Flask, render_template, Response, jsonify, request, send_file
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import base64
import threading
import time

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024  # 16 MB max for frame uploads

# ── Paths ──────────────────────────────────────────────────────────────────────
BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
DB_PATH    = os.path.join(BASE_DIR, "attendance.db")
FACES_DIR  = os.path.join(BASE_DIR, "faces_db")
REPORTS_DIR = os.path.join(BASE_DIR, "reports")
ENCODINGS_PATH = os.path.join(BASE_DIR, "encodings.pkl")

os.makedirs(FACES_DIR, exist_ok=True)
os.makedirs(REPORTS_DIR, exist_ok=True)

# ── Database ────────────────────────────────────────────────────────────────────
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS students (
            id       INTEGER PRIMARY KEY AUTOINCREMENT,
            name     TEXT NOT NULL,
            roll_no  TEXT UNIQUE NOT NULL,
            dept     TEXT,
            created  TEXT DEFAULT (datetime('now','localtime'))
        );
        CREATE TABLE IF NOT EXISTS attendance (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id INTEGER NOT NULL,
            date       TEXT NOT NULL,
            time_in    TEXT NOT NULL,
            status     TEXT DEFAULT 'Present',
            FOREIGN KEY(student_id) REFERENCES students(id),
            UNIQUE(student_id, date)
        );
    """)
    conn.commit()
    conn.close()

init_db()

# ── Face Recognition State ──────────────────────────────────────────────────────
known_encodings = []
known_ids       = []
recognition_lock = threading.Lock()

ENCODING_VERSION = 2  # bump when feature format changes

def load_encodings():
    global known_encodings, known_ids
    if os.path.exists(ENCODINGS_PATH):
        with open(ENCODINGS_PATH, "rb") as f:
            data = pickle.load(f)
        if data.get("version") == ENCODING_VERSION:
            known_encodings = data.get("encodings", [])
            known_ids       = data.get("ids", [])
        else:
            print(f"[Face Recognition] Old encoding format (v{data.get('version', 0)}) discarded — please re-register faces.")
            known_encodings = []
            known_ids       = []
    else:
        known_encodings = []
        known_ids       = []

def save_encodings():
    with open(ENCODINGS_PATH, "wb") as f:
        pickle.dump({"version": ENCODING_VERSION, "encodings": known_encodings, "ids": known_ids}, f)

load_encodings()

# ── Camera & Recognition Thread ─────────────────────────────────────────────────
camera      = None
camera_lock = threading.Lock()
last_frame  = None
last_recognized = {}   # student_id -> timestamp
COOLDOWN    = 5        # seconds between re-marking same person

try:
    face_cascade = cv2.CascadeClassifier(
        cv2.data.haarcascades + "haarcascade_frontalface_default.xml"
    )
    eye_cascade = cv2.CascadeClassifier(
        cv2.data.haarcascades + "haarcascade_eye.xml"
    )
except Exception:
    face_cascade = None
    eye_cascade = None

def align_face(face_img):
    """Align face based on eye positions for consistent encoding."""
    if cv2 is None:
        return face_img
    try:
        gray = cv2.cvtColor(face_img, cv2.COLOR_BGR2GRAY) if len(face_img.shape) == 3 else face_img
        h, w = gray.shape[:2]

        eyes = []
        if eye_cascade is not None:
            eyes = eye_cascade.detectMultiScale(gray, 1.1, 5, minSize=(10, 10))

        if len(eyes) >= 2:
            # Sort eyes by x-coordinate to get left and right eye
            eyes = sorted(eyes, key=lambda e: int(e[0]))
            left_eye  = eyes[0]
            right_eye = eyes[1]
            # Centers (cast to Python int for OpenCV compatibility)
            lx = int(left_eye[0])  + int(left_eye[2])  // 2
            ly = int(left_eye[1])  + int(left_eye[3])  // 2
            rx = int(right_eye[0]) + int(right_eye[2]) // 2
            ry = int(right_eye[1]) + int(right_eye[3]) // 2
            # Angle
            angle = float(np.degrees(np.arctan2(ry - ly, rx - lx)))
            # Center between eyes
            cx, cy = (lx + rx) // 2, (ly + ry) // 2
            # Rotation matrix
            M = cv2.getRotationMatrix2D((cx, cy), angle, 1.0)
            face_img = cv2.warpAffine(face_img, M, (w, h))
    except Exception:
        pass  # if alignment fails, return original face
    return face_img

def get_camera():
    global camera
    with camera_lock:
        if cv2 is None:
            camera = None
            return None
        if camera is None or not camera.isOpened():
            camera = cv2.VideoCapture(0)
            camera.set(cv2.CAP_PROP_FRAME_WIDTH,  640)
            camera.set(cv2.CAP_PROP_FRAME_HEIGHT, 480)
    return camera

def release_camera():
    global camera
    with camera_lock:
        if camera and camera.isOpened():
            camera.release()
        camera = None

def compute_encoding(face_img):
    """Discriminative encoding using face alignment + HOG + histogram features."""
    if cv2 is None:
        return np.zeros(100, dtype=np.float32)

    # Align face using eye positions
    face_img = align_face(face_img)

    # Convert to grayscale
    gray = cv2.cvtColor(face_img, cv2.COLOR_BGR2GRAY) if len(face_img.shape) == 3 else face_img

    # Resize to standard size
    face_size = 64
    gray = cv2.resize(gray, (face_size, face_size))

    # Apply CLAHE for illumination normalization
    clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
    gray = clahe.apply(gray)

    # --- HOG features (captures edge/shape structure) ---
    hog = cv2.HOGDescriptor(
        _winSize=(face_size, face_size),
        _blockSize=(16, 16),
        _blockStride=(8, 8),
        _cellSize=(8, 8),
        _nbins=9,
        _derivAperture=1,
        _winSigma=-1,
        _histogramNormType=cv2.HOGDescriptor_L2Hys,
        _L2HysThreshold=0.2,
        _gammaCorrection=True,
        _nlevels=cv2.HOGDescriptor_DEFAULT_NLEVELS
    )
    hog_features = hog.compute(gray).flatten().astype(np.float32)

    # Normalize HOG
    norm = np.linalg.norm(hog_features)
    if norm > 0:
        hog_features = hog_features / norm

    # --- Grayscale histogram (captures intensity distribution) ---
    hist = cv2.calcHist([gray], [0], None, [64], [0, 256]).flatten().astype(np.float32)
    norm = np.linalg.norm(hist)
    if norm > 0:
        hist = hist / norm

    # Combine features with histogram weighted lower
    features = np.concatenate([hog_features, hist * 0.5])
    return features

def compare_encodings(known_encs, candidate, threshold=0.60):
    """Group-based mean distance comparison with ratio check.

    Instead of simple vote counting, this:
    1. Groups all samples by student
    2. Computes mean of best-K distances per student (more robust)
    3. Requires the best match to be significantly better than the runner-up
    """
    if not known_encs:
        return -1, 1.0

    candidate = np.array(candidate, dtype=np.float32)

    # Compute distances from candidate to every stored encoding
    diffs = []
    for e in known_encs:
        enc = np.array(e, dtype=np.float32)
        if enc.shape != candidate.shape:
            diffs.append(1.0)  # dimension mismatch → no match
        else:
            diffs.append(float(np.linalg.norm(enc - candidate)))

    # Group distances by student_id
    student_dists = {}
    for i, d in enumerate(diffs):
        sid = known_ids[i]
        student_dists.setdefault(sid, []).append(d)

    if not student_dists:
        return -1, 1.0

    # Compute mean of the K closest samples per student (robust aggregate)
    K = min(5, min(len(v) for v in student_dists.values()) or 1)
    student_mean = {}
    for sid, dists in student_dists.items():
        top_k = sorted(dists)[:K]
        student_mean[sid] = sum(top_k) / len(top_k)

    # Rank students by mean distance
    ranked = sorted(student_mean.items(), key=lambda x: x[1])
    best_sid, best_mean = ranked[0]

    # Must be below threshold
    if best_mean >= threshold:
        return -1, best_mean

    # Ratio check: best must be clearly better than second-best
    if len(ranked) > 1:
        second_mean = ranked[1][1]
        if second_mean > 0 and (best_mean / second_mean) > 0.90:
            return -1, best_mean  # too ambiguous

    best_idx = min(
        (i for i, sid in enumerate(known_ids) if sid == best_sid),
        key=lambda i: diffs[i]
    )
    return best_idx, best_mean

def mark_attendance(student_id):
    now  = datetime.now()
    today = now.strftime("%Y-%m-%d")
    t     = now.strftime("%H:%M:%S")
    conn  = get_db()
    try:
        conn.execute(
            "INSERT OR IGNORE INTO attendance (student_id, date, time_in) VALUES (?,?,?)",
            (student_id, today, t)
        )
        conn.commit()
    finally:
        conn.close()

def gen_frames():
    global last_frame, last_recognized
    cam = get_camera()
    if cam is None:
        return
    while True:
        try:
            ok, frame = cam.read()
            if not ok:
                time.sleep(0.05)
                continue

            gray  = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            faces = []
            if face_cascade is not None:
                faces = face_cascade.detectMultiScale(gray, 1.1, 5, minSize=(80, 80))

            for (x, y, w, h) in faces:
                try:
                    face_crop = frame[y:y+h, x:x+w]
                    enc       = compute_encoding(face_crop)

                    with recognition_lock:
                        idx, dist = compare_encodings(known_encodings, enc)

                    label = "Unknown"
                    color = (0, 0, 220)

                    if idx != -1:
                        sid = known_ids[idx]
                        conn = get_db()
                        row  = conn.execute("SELECT name, roll_no FROM students WHERE id=?", (sid,)).fetchone()
                        conn.close()
                        if row:
                            label = f"{row['name']} ({row['roll_no']}) [{dist:.3f}]"
                            color = (0, 200, 60)
                            now_ts = time.time()
                            if sid not in last_recognized or (now_ts - last_recognized[sid]) > COOLDOWN:
                                last_recognized[sid] = now_ts
                                mark_attendance(sid)
                    else:
                        # Show nearest distance for debugging
                        label = f"Unknown ({dist:.3f})"
                except Exception:
                    label = "Unknown"
                    color = (0, 0, 220)

                cv2.rectangle(frame, (x, y), (x+w, y+h), color, 2)
                cv2.rectangle(frame, (x, y-30), (x+w, y), color, -1)
                cv2.putText(frame, label, (x+4, y-8),
                            cv2.FONT_HERSHEY_SIMPLEX, 0.55, (255,255,255), 1)

            # Timestamp overlay
            ts = datetime.now().strftime("%d-%m-%Y  %H:%M:%S")
            cv2.putText(frame, ts, (10, frame.shape[0]-10),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.5, (200,200,200), 1)

            last_frame = frame.copy()
            ret, buf = cv2.imencode(".jpg", frame, [cv2.IMWRITE_JPEG_QUALITY, 80])
            if ret:
                yield (b"--frame\r\nContent-Type: image/jpeg\r\n\r\n" + buf.tobytes() + b"\r\n")
        except Exception:
            time.sleep(0.05)
            continue

# ── Routes ─────────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    return render_template("index.html")

@app.route("/video_feed")
def video_feed():
    if cv2 is None:
        return Response("Camera unavailable in this environment.", mimetype="text/plain")
    return Response(gen_frames(), mimetype="multipart/x-mixed-replace; boundary=frame")

@app.route("/api/process-frame", methods=["POST"])
def process_frame():
    """Process a frame sent from the browser's webcam.

    Accepts a base64-encoded JPEG, runs face detection + recognition,
    and returns face bounding boxes with labels.
    """
    if cv2 is None or face_cascade is None:
        return jsonify({"faces": []})

    try:
        data = request.get_json()
        image_data = data.get("image", "")
        # Strip data URL prefix if present
        if "," in image_data:
            image_data = image_data.split(",", 1)[1]

        img_bytes = base64.b64decode(image_data)
        nparr = np.frombuffer(img_bytes, np.uint8)
        frame = cv2.imdecode(nparr, cv2.IMREAD_COLOR)

        if frame is None:
            return jsonify({"faces": []})

        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        faces = face_cascade.detectMultiScale(gray, 1.1, 5, minSize=(80, 80))

        results = []
        for (x, y, w, h) in faces:
            face_crop = frame[y:y+h, x:x+w]
            enc = compute_encoding(face_crop)

            with recognition_lock:
                idx, dist = compare_encodings(known_encodings, enc)

            face_info = {
                "x": int(x), "y": int(y), "w": int(w), "h": int(h),
                "label": "Unknown",
                "distance": round(float(dist), 3),
                "recognized": False,
                "student_id": None,
            }

            if idx != -1:
                sid = known_ids[idx]
                conn = get_db()
                row = conn.execute("SELECT name, roll_no FROM students WHERE id=?", (sid,)).fetchone()
                conn.close()
                if row:
                    face_info["label"] = f"{row['name']} ({row['roll_no']})"
                    face_info["recognized"] = True
                    face_info["student_id"] = sid

                    now_ts = time.time()
                    if sid not in last_recognized or (now_ts - last_recognized[sid]) > COOLDOWN:
                        last_recognized[sid] = now_ts
                        mark_attendance(sid)

            results.append(face_info)

        return jsonify({"faces": results})
    except Exception as e:
        return jsonify({"faces": [], "error": str(e)})

@app.route("/api/students", methods=["GET"])
def get_students():
    conn = get_db()
    rows = conn.execute("SELECT * FROM students ORDER BY name").fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])

@app.route("/api/students", methods=["POST"])
def add_student():
    data    = request.json
    name    = data.get("name","").strip()
    roll_no = data.get("roll_no","").strip()
    dept    = data.get("dept","").strip()
    if not name or not roll_no:
        return jsonify({"error": "Name and Roll No are required"}), 400
    conn = get_db()
    try:
        cur = conn.execute(
            "INSERT INTO students (name, roll_no, dept) VALUES (?,?,?)",
            (name, roll_no, dept)
        )
        conn.commit()
        sid = cur.lastrowid
        return jsonify({"success": True, "id": sid})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Roll number already exists"}), 409
    finally:
        conn.close()

@app.route("/api/capture", methods=["POST"])
def capture_face():
    """Capture face samples from browser-sent frames for a student."""
    data = request.json
    sid  = data.get("student_id")
    if sid is None:
        return jsonify({"error": "student_id required"}), 400

    if cv2 is None or face_cascade is None:
        return jsonify({"error": "OpenCV is unavailable in this environment."}), 400

    frames = data.get("frames", [])
    if not frames:
        return jsonify({"error": "No frames received. Ensure camera is active."}), 400

    samples = []
    for frame_data in frames:
        if "," in frame_data:
            frame_data = frame_data.split(",", 1)[1]
        try:
            img_bytes = base64.b64decode(frame_data)
            nparr = np.frombuffer(img_bytes, np.uint8)
            frame = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
            if frame is None:
                continue
            gray  = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            faces = face_cascade.detectMultiScale(gray, 1.1, 5, minSize=(80, 80))
            for (x, y, w, h) in faces:
                face_crop = frame[y:y+h, x:x+w]
                enc = compute_encoding(face_crop)
                samples.append(enc)
                if len(samples) >= 20:
                    break
        except Exception:
            continue
        if len(samples) >= 20:
            break

    if len(samples) < 5:
        return jsonify({"error": "Could not capture enough face samples. Ensure face is visible and well-lit."}), 400

    with recognition_lock:
        # Remove existing encodings for this student
        new_encs = []
        new_ids  = []
        for e, i in zip(known_encodings, known_ids):
            if i != sid:
                new_encs.append(e)
                new_ids.append(i)
        for s in samples:
            new_encs.append(s.tolist())
            new_ids.append(sid)
        known_encodings.clear(); known_encodings.extend(new_encs)
        known_ids.clear();       known_ids.extend(new_ids)
        save_encodings()

    return jsonify({"success": True, "samples": len(samples)})

@app.route("/api/attendance", methods=["GET"])
def get_attendance():
    target = request.args.get("date", date.today().isoformat())
    conn   = get_db()
    rows   = conn.execute("""
        SELECT s.name, s.roll_no, s.dept, a.time_in, a.status
        FROM attendance a
        JOIN students s ON a.student_id = s.id
        WHERE a.date = ?
        ORDER BY a.time_in
    """, (target,)).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])

@app.route("/api/attendance/stats", methods=["GET"])
def attendance_stats():
    today = date.today().isoformat()
    conn  = get_db()
    total   = conn.execute("SELECT COUNT(*) FROM students").fetchone()[0]
    present = conn.execute(
        "SELECT COUNT(*) FROM attendance WHERE date=?", (today,)
    ).fetchone()[0]
    conn.close()
    return jsonify({"total": total, "present": present, "absent": total - present, "date": today})

@app.route("/api/export", methods=["GET"])
def export_excel():
    target = request.args.get("date", date.today().isoformat())
    conn   = get_db()
    rows   = conn.execute("""
        SELECT s.name, s.roll_no, s.dept, a.time_in, a.status
        FROM attendance a JOIN students s ON a.student_id = s.id
        WHERE a.date=? ORDER BY a.time_in
    """, (target,)).fetchall()
    all_students = conn.execute("SELECT name, roll_no, dept FROM students ORDER BY name").fetchall()
    conn.close()

    present_rolls = {r["roll_no"] for r in rows}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Attendance {target}"

    # ── Styles
    hdr_fill  = PatternFill("solid", fgColor="1A237E")
    hdr_font  = Font(color="FFFFFF", bold=True, size=11)
    alt_fill  = PatternFill("solid", fgColor="E8EAF6")
    pres_fill = PatternFill("solid", fgColor="C8E6C9")
    abs_fill  = PatternFill("solid", fgColor="FFCDD2")
    center    = Alignment(horizontal="center", vertical="center")
    thin      = Side(style="thin", color="BDBDBD")
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title
    ws.merge_cells("A1:E1")
    ws["A1"] = "ATTENDANCE REPORT"
    ws["A1"].font = Font(bold=True, size=14, color="1A237E")
    ws["A1"].alignment = center

    ws.merge_cells("A2:E2")
    ws["A2"] = f"Date: {target}   |   Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    ws["A2"].alignment = center
    ws["A2"].font = Font(italic=True, color="555555")

    ws.append([])  # blank row

    headers = ["#", "Name", "Roll No", "Department", "Status / Time In"]
    ws.append(headers)
    hdr_row = ws.max_row
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(hdr_row, col)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = center
        cell.border = border

    present_map = {r["roll_no"]: r for r in rows}
    for i, stu in enumerate(all_students, 1):
        if stu["roll_no"] in present_map:
            r = present_map[stu["roll_no"]]
            row_data = [i, stu["name"], stu["roll_no"], stu["dept"] or "—", f"Present  {r['time_in']}"]
            fill = pres_fill
        else:
            row_data = [i, stu["name"], stu["roll_no"], stu["dept"] or "—", "Absent"]
            fill = abs_fill

        ws.append(row_data)
        cur_row = ws.max_row
        for col in range(1, 6):
            cell = ws.cell(cur_row, col)
            cell.fill = fill if col == 5 else (alt_fill if i % 2 == 0 else PatternFill())
            cell.alignment = center
            cell.border = border

    # Column widths
    for col, width in zip("ABCDE", [5, 22, 14, 18, 22]):
        ws.column_dimensions[col].width = width

    # Summary
    ws.append([])
    ws.append(["", "Total Students", len(all_students), "", ""])
    ws.append(["", "Present", len(present_rolls), "", ""])
    ws.append(["", "Absent", len(all_students) - len(present_rolls), "", ""])

    path = os.path.join(REPORTS_DIR, f"attendance_{target}.xlsx")
    wb.save(path)
    return send_file(path, as_attachment=True,
                     download_name=f"attendance_{target}.xlsx")

@app.route("/api/students/<int:sid>", methods=["DELETE"])
def delete_student(sid):
    conn = get_db()
    conn.execute("DELETE FROM attendance WHERE student_id=?", (sid,))
    conn.execute("DELETE FROM students WHERE id=?", (sid,))
    conn.commit()
    conn.close()
    with recognition_lock:
        indices = [i for i,x in enumerate(known_ids) if x != sid]
        new_encs = [known_encodings[i] for i in indices]
        new_ids  = [known_ids[i] for i in indices]
        known_encodings.clear(); known_encodings.extend(new_encs)
        known_ids.clear();       known_ids.extend(new_ids)
        save_encodings()
    return jsonify({"success": True})

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    print(f"[Face Recognition] Attendance System running -> http://localhost:{port}")
    app.run(debug=False, threaded=True, host="0.0.0.0", port=port)
